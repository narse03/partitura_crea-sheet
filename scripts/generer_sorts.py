#!/usr/bin/env python3
"""Génère src/data/sorts.json à partir du Livret des Voix (docx).

Usage : python3 scripts/generer_sorts.py chemin/vers/livret_des_voix.docx

Seules les sept Voix jouables sont exportées ; les Voix réservées au Meneur
(Éléments, Contrôle, Fin et Renouveau) restent hors de l'application.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document

VOIX = {
    'VOIX UNIVERSELLE': 'universelle',
    'VOIX DES ARMES': 'armes',
    'VOIX SAUVAGE': 'sauvage',
    'VOIX DES DIEUX': 'dieux',
    'VOIX DES OMBRES': 'ombres',
    'VOIX DES ÉRUDITS': 'erudits',
    'VOIX DE LA CRÉATION': 'creation',
}
# Ordre des paragraphes « Les Paliers de la Voix » dans le Livret
ORDRE_TITRES = ['universelle', 'armes', 'sauvage', 'dieux', 'ombres', 'erudits', 'creation']
CATS = ['mineurs', 'utilitaires', 'tactiques', 'signature', 'rituels']
TYPES = {'Standard': 'S', 'Instantané': 'I', 'Concentration': 'C', 'Réaction': 'X'}
PALIERS = ['Novice', 'Pratiquant', 'Expert', 'Maître', 'Magistère']


def clean(s: str) -> str:
    return re.sub(r'[ \t]+', ' ', s.replace('\n', ' ')).strip()


def num(s: str):
    s = clean(s)
    return int(s) if s.isdigit() else s


def main(src: str) -> None:
    doc = Document(src)

    titres = []
    for p in doc.paragraphs:
        if p.text.startswith('Les Paliers de la Voix'):
            t = {}
            for nom, pal in re.findall(r'([^,:;]+?) \((Novice|Pratiquant|Expert|Maître|Magistère)\)', p.text):
                nom = re.sub(r'^(et )', '', nom.strip())
                t[pal] = nom[0].upper() + nom[1:]
            titres.append(t)

    data = {v: {'titres': titres[i] if i < len(titres) else {}, 'sorts': {c: [] for c in CATS}}
            for i, v in enumerate(ORDRE_TITRES)}

    cur, k = None, 0
    for t in doc.tables:
        first = t.rows[0].cells[0].text.strip()
        if len(t.rows) == 1 and len(t.columns) == 1 and first.upper().startswith('VOIX'):
            cur, k = VOIX.get(first.split('\n')[0].strip()), 0
            continue
        head = [c.text.strip() for c in t.rows[0].cells]
        if cur is None or head[:2] not in (['Sort', 'Diff'], ['Rituel', 'Diff']):
            continue
        cat = CATS[k]
        k += 1
        for r in t.rows[1:]:
            c = [clean(x.text) for x in r.cells]
            if head[0] == 'Rituel':
                s = {'n': c[0], 'diff': num(c[1]), 'pm': num(c[2]), 'pv': num(c[3]),
                     'type': 'R', 'incant': c[4], 'dur': c[5], 'effet': c[6]}
            else:
                typ = c[3]
                code = TYPES.get(typ)
                s = {'n': c[0], 'diff': num(c[1]), 'pm': num(c[2]), 'type': code or 'S',
                     'dur': c[4], 'effet': c[5]}
                if code is None:
                    # Type hors norme (ex. « Rituel court (30 min) ») : texte conservé tel quel
                    s['type'] = 'R' if typ.lower().startswith('rituel') else 'S'
                    s['typeTexte'] = typ
            data[cur]['sorts'][cat].append(s)

    out = Path(__file__).resolve().parent.parent / 'src' / 'data' / 'sorts.json'
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding='utf-8')
    total = sum(len(l) for v in data.values() for l in v['sorts'].values())
    print(f'{total} sorts écrits dans {out}')


if __name__ == '__main__':
    main(sys.argv[1])
